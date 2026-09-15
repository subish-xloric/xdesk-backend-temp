from types import SimpleNamespace
from datetime import datetime, timedelta
from unittest import result

from django.http import HttpResponse
#from httplib2 import Response
# from django.conf import settings

from pTracker.common.utility import Utility
from pTracker.common.exception_handler import ExceptionHandler
from pTracker.common.logs import Logs
from pTracker.common.file_manager import FileManager

from pTracker.cronjobs.email_sender import send_email_notification
from pTracker.dataaccess.ptracker_access.appraisal_da import AppraisalDA
from pTracker.dataaccess.ptracker_access.user_da import UserDA

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Color, Fill
from openpyxl.styles import Font

import json

from pTracker import settings


def new_dto():
    dto = SimpleNamespace()
    return dto


class AppraisalReportBL():

    def __init__(self):
        self.__log = Logs()
        self.__exception = ExceptionHandler()
        self.__utility = Utility()
        self.__file_manager = FileManager()

    def __get_clean_integer(self, app_year):
        try:
            return int(app_year)
        except:
            return None
    def __get_organization_by_org_id(self, organization_id):
        organization = ''
        if organization_id == 2:
            organization = 'Digitalmesh'
        else:
            organization = "EM Softech"
        return  organization

    def __get_active_emp_dict(self):
        emp_dict = {}
        users = UserDA().get_all_active_users()
        for user in users:
            emp_dict[user.id] = user
        return emp_dict

    def __get_emp_name(self, emp):
        if emp:
            return emp.first_name + ' ' + emp.last_name
        else:
            return ''

    def __get_all_emp_dict(self):
        emp_dict = {}
        users = UserDA().get_all_users()
        for user in users:
            emp_dict[user.id] = user
        return emp_dict

    def __user_profile_dict(self):
        profile_dict = {}
        profiles = UserDA().get_all_user_profiles()
        for profile in profiles:
            profile_dict[profile.user_id] = profile #int(profile.company_id)
        return profile_dict
    def _get_job_title_dict(self):
        job_title_dict = {}
        job_titles = UserDA().get_all_job_titles()
        for job in job_titles:
            job_title_dict[job.id] = job.job_title
        return job_title_dict

    def __get_years_of_experience(self, date_joined):
        try:
            months_of_service = datetime.now().month\
                - date_joined.month + 12 * \
                    (datetime.now().year - date_joined.year)
            years, months = divmod(months_of_service, 12)
            return str(years) + ' Years ' + str(months) + ' Months'
        except:
            return "--"

    def get_appraisal_period_range_years(self, appraisal_period_obj):
        start_year  = appraisal_period_obj.period_start_date.year
        end_year = appraisal_period_obj.period_end_date.year
        return start_year,end_year



    def generate_appraisal_excel_report(self, user_id, year, organization, batch_id):
        result = {"error": '', 'status': 200, "data": [] }
        try:
            role_id, role_name = UserDA().get_user_role_by_id(user_id)
            if role_id not in [1, 2] :
                result['error'] = settings.ERROR_MSG['access_denied']
                result['status'] = 403
                return result

            batch_id = self.__get_clean_integer(batch_id)
            batch = AppraisalDA().get_appraisal_batch_by_id(batch_id)

            if not batch:
                result['error'] = "Invalid batch provided."
                result['status'] = 499
                return result

            year = self.__get_clean_integer(year)
            if not year:
                result['error'] = "Invalid year provided."
                result['status'] = 499
                return result

            appraisal_period = AppraisalDA().get_appraisal_period_by_date(year)
            if not appraisal_period:
                result['error'] = "No appraisal period recoreded in system for the year " + str(year)
                result['status'] = 499
                return result

            organization = self.__get_clean_integer(organization)

            forms = AppraisalDA().get_all_my_appraisal_forms_by_appraisal_period_and_batch_organization(appraisal_period.period_id, batch[0].batch_id, organization)
            output =self.format_to_excel(organization, batch, appraisal_period, forms)
        except Exception as err:
            output = ''
            result["error"] = settings.ERROR_MSG['application_error']\
                .format(err, self.__log.error(self.__exception.get_exception()))
        return output

    def format_to_excel(self, organization, batch, appraisal_period, appraisal_forms):
        output = HttpResponse(
                content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet;charset=utf-8;'
            )
        file_name = "AppraisalReport.xlsx"
        output['Content-Disposition'] = 'attachment; filename=' + file_name

        company_name = self.__get_organization_by_org_id(organization)
        users = self.__get_all_emp_dict()
        user_profile_dict = self.__user_profile_dict()
        job_title_dict = self._get_job_title_dict()
        appraisal_start_date_year, appraisal_end_date_year = self.get_appraisal_period_range_years(appraisal_period)

        bold = Font(bold=True)
        header_fill = PatternFill(
            start_color="F0F802", end_color="F0F802", fill_type='solid')

        if appraisal_forms:
            wb = openpyxl.Workbook()
            ws = wb.worksheets[0]
            if company_name == 'Digitalmesh':
                company_address = company_address = """Digital Mesh Softech India (P) Limited\nUnit 1: 43-A, E Block, 2nd Floor,\nCochin Special Economic Zone, Kakkanad, Kochi – 682 037, Kerala, India.\nTel: +91-484-4060200, Fax: +91-484-4060201"""
            else:
                company_address = """EM Softech LLP\nUnit 1:Plot No.43/ A, D Block, 2nd floor,\nCochin Special Economic Zone(CSEZ), Kakkanad, Kochi-682037, Kerala, India.\nTel:+91-484-2413280"""
            day_head = ['R']
            ws.merge_cells('A1:' + day_head[-1]+'1')
            ws.merge_cells('A3:' + day_head[-1]+'3')

            ws['A1'].value = company_address
            ws['A1'].font = bold
            ws['A1'].alignment = Alignment(horizontal='center')
            ws.row_dimensions[1].height = 95

            ws['A3'].value = 'Performance Assessment Report of ' + str(appraisal_start_date_year)+'-'+ str(appraisal_end_date_year)+ '('+  \
                batch[0].batch_name+" )"
            ws['A3'].font = bold
            ws['A3'].alignment = Alignment(horizontal='center')
            ws['A3'].fill = header_fill

            ws['A5'].value = 'Sl No.'
            ws['A5'].font = bold

            ws['B5'].value = 'Emp Code'
            ws['B5'].font = bold

            ws['C5'].value = 'Emp Name'
            ws['C5'].font = bold

            ws['D5'].value = 'DOJ'
            ws['D5'].font = bold

            ws['E5'].value = 'Designation'
            ws['E5'].font = bold

            ws['F5'].value = 'Pre-DM Exp'
            ws['F5'].font = bold

            ws['G5'].value = 'In DM exp'
            ws['G5'].font = bold

            ws['H5'].value = 'Total yrs of Exp'
            ws['H5'].font = bold

            ws['I5'].value = 'Rating'
            ws['I5'].font = bold

            ws['J5'].value = 'Rating %'
            ws['J5'].font = bold

            ws['K5'].value = 'CTC Per Month'
            ws['K5'].font = bold

            ws['L5'].value = 'Proposed Increment'
            ws['L5'].font = bold

            ws['M5'].value = 'Proposed Salary'
            ws['M5'].font = bold

            ws['N5'].value = 'Final Salary'
            ws['N5'].font = bold

            ws['O5'].value = 'Remarks & Recommendation to Management'
            ws['O5'].font = bold

            ws['P5'].value = 'Project Handled'
            ws['P5'].font = bold

            ws['Q5'].value = 'PF/ESI'
            ws['Q5'].font = bold

            row = 7
            counter = 1
            for each_form in appraisal_forms:

                user = users.get(each_form.employee_id, None)
                user_profile = user_profile_dict.get(
                    each_form.employee_id, None)
                job_title = job_title_dict.get(
                    int(user_profile.job_title), None)
                experience = self.__get_years_of_experience(user.date_joined)
                appraisal_rating = ''
                rating_percentage = ''

                if each_form.rating:
                    appraisal_rating = settings.PERSONAL_APPRAISAL_RATINGS[each_form.rating]
                    rating_percentage = settings.APPRAISAL_RATING_PERCENTAGE[each_form.rating]
                if user_profile:
                    job_title_id = user_profile.job_title

                appraisal_form = json.loads(each_form.appraisal_data)
                if user:
                    emp_name = user.first_name + ' ' + user.last_name
                    ws['A'+str(row)].value = counter
                    ws['A'+str(row)].alignment = Alignment(horizontal='center')
                    ws['B'+str(row)].value = user.username
                    ws['C'+str(row)].value = emp_name
                    ws['D'+str(row)].value = user.date_joined.strftime("%m/%d/%Y")
                    ws['E'+str(row)].value = job_title
                    ws['G'+str(row)].value = experience
                    ws['H'+str(row)].value = experience
                    ws['I'+str(row)].value = appraisal_rating
                    ws['J'+str(row)].value = rating_percentage
                    ws['O'+str(row)].value = appraisal_form['recommendation_to_management']['value']
                    ws['O'+str(row)].alignment = Alignment(wrap_text=True)
                    ws['P'+str(row)].alignment = Alignment(wrap_text=True)
                row += 1
                counter += 1

            ws.column_dimensions["A"].width = 6
            ws.column_dimensions["B"].width = 14
            ws.column_dimensions["C"].width = 25
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 35
            ws.column_dimensions["F"].width = 25
            ws.column_dimensions["G"].width = 25
            ws.column_dimensions["H"].width = 25
            ws.column_dimensions["I"].width = 25
            ws.column_dimensions["J"].width = 10
            ws.column_dimensions["K"].width = 25
            ws.column_dimensions["L"].width = 25
            ws.column_dimensions["M"].width = 25
            ws.column_dimensions["N"].width = 25
            ws.column_dimensions["O"].width = 50
            ws.column_dimensions["P"].width = 40
            ws.column_dimensions["Q"].width = 25
            wb.save(output)
        return output


    def get_appraisal_response_report(self, user_id, year,  batch_id , response_of):
        result = {"error": '', "status": 200, "data": [], "is_training_selected": False}
        try:
            role_id, role_name = UserDA().get_user_role_by_id(user_id)
            if role_id not in [1, 2] :
                result['error'] = settings.ERROR_MSG['access_denied']
                result['status'] = 403
                return result

            batch_id = self.__get_clean_integer(batch_id)
            batch = AppraisalDA().get_appraisal_batch_by_id(batch_id)

            if not batch:
                result['error'] = "Invalid batch provided."
                result['status'] = 499
                return result

            year = self.__get_clean_integer(year)
            if not year:
                result['error'] = "Invalid year provided."
                result['status'] = 499
                return result

            appraisal_period = AppraisalDA().get_appraisal_period_by_date(year)
            if not appraisal_period:
                result['error'] = "No appraisal period recoreded in system for the year " + str(year)
                result['status'] = 499
                return result
            response_of = self.__get_clean_integer(response_of)
            if not response_of:
                result['error'] = "Invalid Response type."
                result['status'] = 499
                return result

            # organization = self.__get_clean_integer(organization)

            appraisal_forms = AppraisalDA().get_all_my_appraisal_forms_by_appraisal_period_and_batch_organization(appraisal_period.period_id, batch_id)
            report_list, is_training_selected = self.get_formatted_appraisal_respose_report(appraisal_forms, response_of)

            result['data'] = report_list
            result['is_training_selected'] = is_training_selected

        except Exception as err:
            result["error"] = settings.ERROR_MSG['application_error']\
                .format(err, self.__log.error(self.__exception.get_exception()))
        return result

    def get_formatted_appraisal_respose_report(self, appraisal_forms, response_of):
        report_list = []
        is_training_selected = False
        user_dict =  self.__get_all_emp_dict()

        for appraisal_form in appraisal_forms:
            temp = {"response" : '', "technical": '', "non_technical": ''}
            organization = self.__get_organization_by_org_id( appraisal_form.organization_id)
            appraisal_json_data = json.loads(appraisal_form.appraisal_data)
            employee_info = appraisal_json_data['employee']
            self_assesment_data =  appraisal_json_data['self_assessment_form']
            justification_data = appraisal_json_data['justification_for_rating']['value']
            recommendation = appraisal_json_data['recommendation_to_management']['value']
            technical_training_data = appraisal_json_data['training_need_identification']['content_b']['training_programs']

            response_attribute = settings.APPRAISAL_RESPONSE_ATTRIBUTES[response_of]

            if response_attribute:
                section_heading = response_attribute['parent']
                section_attribute = response_attribute['attribute']

                employee = user_dict.get(employee_info['employee_id'], None)
                if employee:

                    temp['emp_name'] = self.__get_emp_name(employee)
                    temp['emp_code'] = employee.username
                    temp['organization'] = organization

                    if section_heading == "self_assessment_form":
                        for qstn in self_assesment_data['qustions']:
                            if qstn['qustion'] == section_attribute:
                                temp['response'] = qstn.get('answer', '')

                    if section_heading == "justification_for_rating":
                        temp['response'] = justification_data

                    if section_heading == "recommendation_to_management":
                        temp['response'] = recommendation

                    if section_heading == "training_need_identification":
                        is_training_selected = True
                        technical = ''

                        non_technical = ''
                        for technical_program in  technical_training_data['technical']:
                            technical = technical+technical_program['value']+','

                        temp['technical'] = technical
                        for non_technical_program in  technical_training_data['non_technical']:
                            non_technical = non_technical+non_technical_program['value']+','
                        temp['non_techical'] = non_technical

                    report_list.append(temp)

        return report_list, is_training_selected

    def upload_appraisal_document(self,user_id, formdata):
        from PyPDF2 import PdfFileWriter, PdfFileReader
        result= {"error": ''}
        try:
            out = PdfFileWriter()
            doc = formdata.FILES['image']
            file = PdfFileReader(doc)

            num = file.numPages
            for idx in range(num):
                # Get the page at index idx
                page = file.getPage(idx)
                # Add it to the output file
                out.addPage(page)
            password = "pass"
            out.encrypt(password)
            import io
            pdf_bytes = io.BytesIO()
            out.write(pdf_bytes)
            
            # Using FileManager instead of hardcoded local path
            file_path = f"{settings.CONFIDENTIAL_DOCS}appraisal_reports/myfile.pdf"
            self.__file_manager.upload_file(file_path, pdf_bytes.getvalue())

            pass
        except Exception as err:
            result["error"] = settings.ERROR_MSG['application_error']\
                .format(err, self.__log.error(self.__exception.get_exception()))
        return result







