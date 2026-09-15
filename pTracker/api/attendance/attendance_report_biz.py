
import os
import time
import calendar
from  datetime import datetime, date, timedelta
from types import SimpleNamespace
from io import BytesIO

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Color, Fill
from openpyxl.styles import Font

from django.conf import settings
from django.http import HttpResponse

from pTracker.api import attendance
from pTracker.dataaccess.db import Connection
from pTracker.common.utility import Utility
from pTracker.common.exception_handler import ExceptionHandler
from pTracker.common.logs import Logs
from pTracker.dataaccess.ptracker_access.project_da import  ProjectDA

from pTracker.dataaccess.essl_access.attendance import  AttendanceDA
from pTracker.dataaccess.ptracker_access.attendance import AttendanceDA as DmAttendanceDA
from pTracker.user_management.holiday_da import HolidayDA
from pTracker.dataaccess.ptracker_access.user_da import UserDA
from pTracker.user_management.employee import Employee


def new_dto():
    dto = SimpleNamespace()
    return dto

def att_dto():
    dto = SimpleNamespace()
    dto.emp_code = None
    dto.emp_name = None
    dto.gender = None
    dto.doj = None
    dto.designation = None
    dto.att_status = 'X'
    return dto

class AttendanceReportBL():

    def __init__(self):
        self.__log = Logs()
        self.__exception = ExceptionHandler()
        self.__utility = Utility()
    
    def __is_weekend(self, date):
        return date.weekday() in [6]  # Saturday is 5, Sunday is 6

    def __is_holiday(self, date):
        is_holiday = HolidayDA().is_holiday(date)
        return is_holiday

    def get_employee_attendence_report(self, request, str_date, organization):
        # str_date :  Should be string with format yyyy-mm-dd
        # TODO- permission

        result = {"error": None, "attendance_report": [], 'company_address': ""}
        att_dict = {}

        try:
            user_id = request.user.id
            is_permitted = self.__check_permission(user_id)
            if not is_permitted:
                result['error'] = settings.ERROR_MSG['access_denied']
                return result

            str_date_obj = datetime.strptime(str_date, "%Y-%m-%d")
            
            if self.__is_holiday(str_date_obj) or self.__is_weekend(str_date_obj):
                return result

            if int(organization) == settings.COMPANY['DM']['ID']:
                result["company_address"] = settings.DM_ADDRESS
                organization = 'DM'
            else:
                result["company_address"] = settings.EM_ADDRESS
                organization = 'EM'

            punch_data = self.__get_punch_data(str_date, organization)
            processed_punch_data = self.__segregate_punch_data(punch_data)

            if processed_punch_data:
                for each in processed_punch_data:
                    att_dict[each.emp_code] = each

            employees = Employee().get_all_active_employees()
            if employees:
                for emp_code in employees:
                    employee = employees[emp_code]
                    dto = att_dto()
                    dto.emp_code = emp_code
                    dto.emp_name = employee.emp_name
                    dto.gender = employee.gender
                    dto.doj = employee.doj
                    dto.designation = employee.designation

                    att_data = att_dict.get(emp_code, None)
                    if not att_data:
                        dto.att_status = 'A'
                    result['attendance_report'].append(dto)
                    del dto
            
        #     result['attendance_report'] =  [
        #   {"slno.":1,"emp_code":574,"emp_name":"Abdul Jaseem P","gender":"Male","doj":"06/09/2021","designation":"Associate Software Engineer","att_status":"X"},
        #   {"slno.":2,"emp_code":572,"emp_name":"Abijith N M","gender":"Male","doj":"06/09/2021","designation":"Associate Software Engineer","att_status":"X"}
        # ]

        except Exception as err:
            result["error"] = settings.ERROR_MSG['application_error']\
                .format(str(err), self.__log.error(self.__exception.get_exception()))
        return result

    
    def __check_permission(self, user_id):
        return True
    

    def __segregate_punch_data(self, punch_data):
        employee_punch_data = {}
        emps = Employee().get_all_active_employees()
        if punch_data:
            for each_item in punch_data:
                dto = new_dto()
                dto.emp_code = str(each_item[0]).strip()
                dto.emp_name = emps.get(dto.emp_code, '-')
                dto.time = Utility().convert_string_to_date_time(each_item[1])
                dto.direction = str(each_item[2]).strip().upper()
                dto.serial_no = str(each_item[3]).strip().upper()

                if dto.emp_code in employee_punch_data:
                    temp_list = employee_punch_data[dto.emp_code]
                    temp_list.append(dto)
                    employee_punch_data[dto.emp_code] = temp_list
                    del temp_list
                else:
                    employee_punch_data[dto.emp_code] = [dto]
                del dto
        return employee_punch_data

    def __get_punch_data(self, str_date, company_code='DM'):
        device_serial_no_in = settings.ATT_DEVICE[company_code + "_IN"]['SerialNumber']
        device_serial_no_out = settings.ATT_DEVICE[company_code + "_OUT"]['SerialNumber']
        str_sql = """SELECT EmployeeCode, LogDate, Direction, DeviceSerialNo
            FROM Att
            WHERE
            (DeviceSerialNo='{0}' OR DeviceSerialNo='{1}')
            AND LogDate >= '{2} 00:00:00' AND LogDate <= '{2} 23:59:59'
            ORDER BY LogDate""".format(device_serial_no_in, device_serial_no_out, str_date)

        #AND EmployeeCode IN ('529','282')
        #print ('str_sql', str_sql)

        # str_sql = """SELECT EmployeeCode, LogDate
        # FROM Att
        # WHERE
        # DeviceSerialNo='{0}' AND EmployeeCode='529'
        # AND LogDate >= '{1} 00:00:00' AND LogDate <= '{1} 23:59:59'
        # ORDER BY LogDate""".format(device_serial_no, str_date)
        conn = Connection('essl_db')
        results, error = conn.execute(str_sql)
        if error:
            Utility().log(error)

        return results


    def get_last_date_of_the_month(self, month, year):
        any_date = datetime(year, month, 1)
        next_month = any_date.replace(day=28) + timedelta(days=4)
        last_day_of_month = next_month - timedelta(days=next_month.day)
        return last_day_of_month

    def get_holidays_for_month(self, start_date, end_date):
        holiday_list = []

        week_ends = Utility().get_all_weekends(start_date, end_date)
        if week_ends:
            for each_day in week_ends:
                try:
                    holiday_list.append(each_day.strftime("%d"))
                except:
                    pass
        holidays = HolidayDA().get_holidays(start_date, end_date)
        for each_day in holidays:
            holiday_list.append(each_day.holiday_date.strftime("%d"))
        return set(holiday_list)
    
    def __procees_emp_attendance(self, att_data):
        day_list = []
        if att_data:
            for each_item in att_data:
                dt = each_item.attendance_date
                day_list.append(dt.day)
        return set(day_list)
    
    def formatDate(self, str_date):
        date_obj = datetime.strptime(str_date, '%Y-%m-%d').strftime("%d/%m/%Y")
        return date_obj
    

    def generate_monthly_att_report_v1(self, request, month, year, organization):
        att_dict = {}
        emp_att_list = []
        result = {'error': '', 'data': [], 'company_address':''}

        #month = "11_2019"

        try:
            user_id = request.user.id
            is_permitted = self.__check_permission(user_id)
            if not is_permitted:
                result['error'] = settings.ERROR_MSG['access_denied']
                return result

            if organization == settings.COMPANY['DM']['ID']:
                result["company_address"] = settings.DM_ADDRESS
            else:
                result["company_address"] = settings.EM_ADDRESS


            start_date = datetime(year, month, 1)
            end_date = self.get_last_date_of_the_month(month, year)

            holidays = []
            temp = self.get_holidays_for_month(start_date, end_date)
            if temp:
                holidays = [int(i) for i in temp]

            start_date = start_date.strftime("%Y-%m-%d")
            end_date = end_date.strftime("%Y-%m-%d")

            days_in_month = calendar.monthrange(year, month)[1]

            employees = Employee().get_all_active_employees()

            att_logs = DmAttendanceDA().get_attendance_by_range(start_date, end_date)
            if att_logs:
                for each_item in att_logs:
                    if each_item.emp_code in att_dict:
                        temp_list = att_dict[each_item.emp_code]
                        temp_list.append(each_item)
                        att_dict[each_item.emp_code] = temp_list

                    else:
                        att_dict[each_item.emp_code] = [each_item]

            if employees:
                for emp_code in employees:
                    employee = employees[emp_code]
                    if employee.company_id!=organization:
                        continue
                    dto = att_dto()
                    dto.emp_code = emp_code
                    dto.emp_name = employee.emp_name
                    dto.gender = employee.gender
                    dto.doj = employee.doj
                    dto.designation = employee.designation
                    dto.day_list = []

                    att_data = att_dict.get(emp_code, None)
                    work_day_list = self.__procees_emp_attendance(att_data)

                    attendance = []
                    for each_day in range(1, 32):
                        try:
                            current_date = datetime.now().date()
                            selected_date = datetime(year, month, each_day).date()

                            if each_day>days_in_month:
                                status='-'
                            elif selected_date > current_date:
                                status = '-'
                            elif each_day in holidays:
                                status = 'H'
                            elif each_day in work_day_list:
                                status = 'X'
                            else:
                                status = 'A'
                        except:
                            status='-'
                        attendance.append({
                                'day': each_day,
                                'status': status
                            })
                    dto.attendance = attendance

                    emp_att_list.append(dto.__dict__)
                    del dto

            result['data'] = emp_att_list

            # self.format_excel_report(dm_emp_att_list, month, year, 'Digitalmesh', holidays)
            # self.format_excel_report(em_emp_att_list, month, year, 'EM_Softech', holidays)

        except Exception as err:
            result["error"] = settings.ERROR_MSG['application_error'] \
                .format(str(err), self.__log.error(self.__exception.get_exception()))
        return result
    

    def generate_monthly_excel_att_report(self, request, month, year, organization):
        att_dict = {}
        emp_att_list = []
        result = None

        #month = "11_2019"

        try:
            user_id = request.user.id
            is_permitted = self.__check_permission(user_id)
            if not is_permitted:
                result['error'] = settings.ERROR_MSG['access_denied']
                return result

            if organization == settings.COMPANY['DM']['ID']:
                company_name = 'Digitalmesh'
            else:
                company_name = 'EM_Softtech'


            start_date = datetime(year, month, 1)
            end_date = self.get_last_date_of_the_month(month, year)

            holidays = []
            temp = self.get_holidays_for_month(start_date, end_date)
            if temp:
                holidays = [int(i) for i in temp]

            start_date = start_date.strftime("%Y-%m-%d")
            end_date = end_date.strftime("%Y-%m-%d")

            days_in_month = calendar.monthrange(year, month)[1]

            employees = Employee().get_all_active_employees()

            att_logs = DmAttendanceDA().get_attendance_by_range(start_date, end_date)
            if att_logs:
                for each_item in att_logs:
                    if each_item.emp_code in att_dict:
                        temp_list = att_dict[each_item.emp_code]
                        temp_list.append(each_item)
                        att_dict[each_item.emp_code] = temp_list

                    else:
                        att_dict[each_item.emp_code] = [each_item]

            if employees:
                for emp_code in employees:
                    employee = employees[emp_code]
                    if employee.company_id!=organization:
                        continue
                    dto = att_dto()
                    dto.emp_code = emp_code
                    dto.emp_name = employee.emp_name
                    dto.gender = employee.gender
                    dto.doj = employee.doj
                    dto.designation = employee.designation
                    dto.day_list = []

                    att_data = att_dict.get(emp_code, None)
                    work_day_list = self.__procees_emp_attendance(att_data)
                    dto.day_list = work_day_list

                    attendance = []
                    for each_day in range(1, 32):
                        try:
                            current_date = datetime.now().date()
                            selected_date = datetime(year, month, each_day).date()

                            if each_day>days_in_month:
                                status='-'
                            elif selected_date > current_date:
                                status = '-'
                            elif each_day in holidays:
                                status = 'H'
                            elif each_day in work_day_list:
                                status = 'X'
                            else:
                                status = 'A'
                        except:
                            status='-'
                        attendance.append({
                                'day': each_day,
                                'status': status
                            })
                    dto.attendance = attendance

                    emp_att_list.append(dto)
                    del dto

            result = self.format_excel_report(emp_att_list, month, year, company_name, holidays)

        except Exception as err:
            result["error"] = settings.ERROR_MSG['application_error'] \
                .format(str(err), self.__log.error(self.__exception.get_exception()))
        return result



    def format_excel_report(self, log_list, month, year, company_name, holidays):

        try:
            output = BytesIO()
            #str_month  = 12_2019
            is_current_month = False

            #month = int(str_month .split("_")[0])
            #year = int(str_month .split("_")[1])

            today = datetime.today()
            current_month = today.month
            current_year = today.year
            current_day = today.day

            if month == current_month and year == current_year:
                is_current_month = True

            #holidays = self.get_holidays_for_month(month, year)
            str_month = datetime(year, month, 1)
            temp_month = str_month.strftime("%B, %Y")
            str_month = str_month.strftime("%B_%Y")


            bold = Font(bold=True)
            header_fill = PatternFill(
                start_color="F0F802", end_color="F0F802", fill_type='solid')

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin'))


            if log_list:
                log_list.sort(key=lambda x: x.emp_name, reverse=False)

                day_head = ['G','H', 'I','J', 'K','L', 'M','N', 'O','P', 'Q','R', 'S','T',
                'U','V', 'W','X', 'Y','Z', 'AA','AB', 'AC','AD', 'AE','AF', 'AG','AH', 'AI','AJ', 'AK']
                remarks_col = "AM"

                if month in [4, 6, 9, 11]:
                    day_head = day_head[:-1]
                    remarks_col = "AL"
                elif month == 2:
                    if self.is_leap_year(year):
                        day_head = day_head[:29]
                        remarks_col = "AK"
                    else:
                        day_head = day_head[:28]
                        remarks_col = "AJ"

                wb = openpyxl.Workbook()
                ws = wb.worksheets[0]


                if company_name == 'Digitalmesh':
                    company_address = """Digital Mesh Softech India (P) Limited\nUnit 1: 43-A, E Block, 2nd Floor,\nCochin Special Economic Zone, Kakkanad, Kochi – 682 037, Kerala, India.\nTel: +91-484-4060200, Fax: +91-484-4060201"""
                else:
                    company_address = """EM Softech LLP\nUnit 1:Plot No.43/ A, D Block, 2nd floor,\nCochin Special Economic Zone(CSEZ), Kakkanad, Kochi-682037, Kerala, India.\nTel:+91-484-2413280"""

                ws.merge_cells('A1:' + day_head[-1]+'1')
                ws.merge_cells('A3:' + day_head[-1]+'3')

                ws['A1'].value = company_address
                ws['A1'].font = bold
                ws['A1'].alignment = Alignment(horizontal='center')
                ws.row_dimensions[1].height = 95


                ws['A3'].value = 'Attendance Report for ' + str(temp_month)
                ws['A3'].font = bold
                ws['A3'].alignment = Alignment(horizontal='center')
                ws['A3'].fill = header_fill

                ws['A5'].value = 'Sl No.'
                ws['A5'].font = bold

                ws['B5'].value = 'Emp Code'
                ws['B5'].font = bold

                ws['C5'].value = 'Emp Name'
                ws['C5'].font = bold

                ws['D5'].value = 'M/F'
                ws['D5'].font = bold

                ws['E5'].value = 'DOJ'
                ws['E5'].font = bold

                ws['F5'].value = 'Designation'
                ws['F5'].font = bold

                head_row = '5'
                index_counter = 0
                for col in day_head:
                    index_counter += 1
                    ws[col+head_row].value = index_counter
                    ws[col+head_row].font = bold
                    ws[col+head_row].alignment = Alignment(horizontal='center')

                row = 7
                counter = 1
                for each_item in log_list:
                    ws['A'+str(row)].value = counter
                    ws['B'+str(row)].value = each_item.emp_code
                    ws['C'+str(row)].value = each_item.emp_name
                    ws['D'+str(row)].value = str(each_item.gender)[0]
                    ws['E'+str(row)].value = self.formatDate(each_item.doj)
                    ws['F'+str(row)].value = each_item.designation


                    """
                    Start date  llop
                    """
                    emp_day_list = each_item.day_list
                    att_day_counter = 1
                    for col in day_head:
                        #ws[col+str(row)].value = each_item.att_status
                        ws[col+str(row)].alignment = Alignment(horizontal='center')
                        #ws[col+str(row)].font = bold

                        if att_day_counter in holidays:
                            ws.column_dimensions[col].width = 3
                            ws[col + str(row)].value = 'H'
                            ws[col + str(row)].font = Font(color="228B22")
                        elif is_current_month and att_day_counter >= current_day:
                            ws[col + str(row)].value = '-'
                            ws.column_dimensions[col].width = 3
                        elif att_day_counter in emp_day_list:
                            ws[col + str(row)].value = 'x'
                            ws.column_dimensions[col].width = 3
                        else:
                            ws[col + str(row)].value = 'A'
                            ws[col + str(row)].font = Font(color="ff0000")
                            ws.column_dimensions[col].width = 3

                        att_day_counter += 1

                    row += 1
                    counter += 1

                ws.column_dimensions["A"].width = 8
                ws.column_dimensions["B"].width = 10
                ws.column_dimensions["C"].width = 20
                ws.column_dimensions["D"].width = 6
                ws.column_dimensions["E"].width = 12
                ws.column_dimensions["F"].width = 25

                temp_file_name = company_name + "_" + str_month + "_original.xlsx"
                wb.save(output)
                output.seek(0)
                 # Create an HTTP response with the Excel file content
                response = HttpResponse(output, content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
                response['Content-Disposition'] = f'attachment; filename="{company_name}_{str_month}.xlsx"'

                return response
        except Exception as err:
            return None
