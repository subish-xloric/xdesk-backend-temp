import os
import datetime
import time
import calendar
#from datetime import timedelta

from django.conf import settings

from types import SimpleNamespace

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Color, Fill
from openpyxl.styles import Font

from pTracker.dataaccess.db import Connection
from pTracker.common.utility import Utility
from pTracker.user_management.employee import Employee
from pTracker.dataaccess.ptracker_access.attendance import AttendanceDA
from pTracker.user_management.holiday_da import HolidayDA
from pTracker.cronjobs.email_sender import send_weekly_report


def att_dto():
    dto = SimpleNamespace()
    dto.emp_code = None
    dto.emp_name = None
    dto.gender = None
    dto.doj = None
    dto.designation = None
    dto.work_hours = 0
    dto.att_status = 'X'
    return dto


def new_dto():
    dto = SimpleNamespace()
    return dto


class WeeklyAttendanceReportBL():

    def get_holidays_for_month(self, start_date, end_date):
        holiday_list = []

        week_ends = Utility().get_all_weekends(start_date, end_date)
        if week_ends:
            for each_day in week_ends:
                try:
                    holiday_list.append(each_day.strftime("%d"))
                except:
                    pass

        start_date = start_date.strftime("%Y-%m-%d")
        end_date = end_date.strftime("%Y-%m-%d")

        obj_holidays = HolidayDA().get_all_holidays(start_date, end_date)
        if obj_holidays:
            for each_item in obj_holidays:
                try:
                    holiday_list.append(each_item.holiday_date.strftime("%d"))
                except Exception as err:
                    pass

        return holiday_list



    def format_excel_report(self, log_list, start_date, end_date):

        bold = Font(bold=True)
        header_fill = PatternFill(
            start_color="F0F802", end_color="F0F802", fill_type='solid')

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'))


        if log_list:
            log_list.sort(key=lambda x: x.emp_name, reverse=False)

            wb = openpyxl.Workbook()
            ws = wb.worksheets[0]
            ws.merge_cells('A1:H1')
            ws['A1'].value = 'Weekly Summary Report from {0} to {1} '.format(start_date, end_date)

            ws['A1'].font = bold
            ws['A1'].alignment = Alignment(horizontal='center')
            ws['A1'].fill = header_fill

            ws['A3'].value = 'Emp Name'
            ws['A3'].font = bold
            #ws['A1'].alignment = Alignment(horizontal='center')

            ws['B3'].value = 'On Time %'
            ws['B3'].font = bold
            ws['B3'].alignment = Alignment(horizontal='center')

            ws['C3'].value = 'Late %'
            ws['C3'].font = bold
            ws['C3'].alignment = Alignment(horizontal='center')

            ws['D3'].value = 'Work Hrs'
            ws['D3'].font = bold
            ws['D3'].alignment = Alignment(horizontal='center')

            ws['E3'].value = 'Hrs Required'
            ws['E3'].font = bold
            ws['E3'].alignment = Alignment(horizontal='center')

            ws['F3'].value = 'Att %'
            ws['F3'].font = bold
            ws['F3'].alignment = Alignment(horizontal='center')

            ws['G3'].value = 'No. Leave'
            ws['G3'].font = bold
            ws['G3'].alignment = Alignment(horizontal='center')


            row = 4
            counter = 1
            for each_item in log_list:
                ws['A'+str(row)].value = each_item.emp_name
                if str(each_item.on_time) == '0.0':
                    on_time = "-"
                else:
                    on_time = each_item.on_time

                ws['B'+str(row)].value = on_time
                ws['B'+str(row)].alignment = Alignment(horizontal='center')

                if str(each_item.late) == '0.0':
                    late = "-"
                else:
                    late = each_item.late

                ws['C'+str(row)].value = late
                ws['C' + str(row)].font = Font(color=each_item.late_color)
                ws['C'+str(row)].alignment = Alignment(horizontal='center')

                ws['D'+str(row)].value = each_item.work_hours
                ws['D' + str(row)].font = Font(color=each_item.work_hours_color)
                ws['D'+str(row)].alignment = Alignment(horizontal='center')

                ws['E'+str(row)].value = each_item.work_hours_required
                ws['E'+str(row)].alignment = Alignment(horizontal='center')
                ws['F'+str(row)].value = each_item.att
                ws['F'+str(row)].alignment = Alignment(horizontal='center')

                if str(each_item.leave) == '0':
                    leave = "-"
                else:
                    leave = each_item.leave

                ws['G'+str(row)].value = leave
                ws['G'+str(row)].alignment = Alignment(horizontal='center')
                row += 1


            ws.column_dimensions["A"].width = 25
            ws.column_dimensions["B"].width = 12
            ws.column_dimensions["C"].width = 12
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 15
            ws.column_dimensions["F"].width = 10
            ws.column_dimensions["G"].width = 12

            upload_path = settings.UPLOAD_PATH['WEEKLY_ATT_REPORT']
            file_name = start_date + "_" + end_date + ".xlsx"
            path_to_file = upload_path + file_name
            wb.save(path_to_file)

            send_weekly_report.apply_async([path_to_file, file_name], queue=settings.CELERY_QUEUE['mail_sender'])


    def __procees_emp_attendance(self, att_data):
        work_hours = 0
        days_punctual = 0
        days_late = 0
        days_present = 0
        #day_list = []
        if att_data:
            for each_item in att_data:
                work_hours += each_item.work_hours
                days_present += 1
                dt = each_item.attendance_date.strftime("%Y-%m-%d")
                punctual_time = settings.PUNCH_IN_CONFIG['punctual']['end_time']
                punctual_time = Utility().create_date_time(str(dt), punctual_time)
                #print ('punctual_time', str(each_item.arrival))
                each_item.arrival = each_item.arrival.strftime("%Y-%m-%d %H:%M:%S")
                arrival = Utility().convert_string_to_date_time(str(each_item.arrival))
                if arrival <= punctual_time:
                    days_punctual += 1
                else:
                    days_late += 1

        dto = new_dto()
        dto.work_hours = work_hours
        dto.days_punctual = days_punctual
        dto.days_late = days_late
        dto.days_present = days_present

        return dto


    def get_last_date_of_the_month(self, month, year):
        any_date = datetime.datetime(year, month, 1)
        next_month = any_date.replace(day=28) + datetime.timedelta(days=4)
        last_day_of_month = next_month - datetime.timedelta(days=next_month.day)
        return last_day_of_month

    def generate_weekly_att_report(self, strat_date, end_date):
        att_dict = {}
        dm_emp_att_list = []
        c_level_emps = []

        #month = "11_2019"

        try:
            start_date = Utility().convert_string_to_date_time(strat_date, "%Y-%m-%d")
            end_date = Utility().convert_string_to_date_time(end_date, "%Y-%m-%d")

            holidays = []
            temp = self.get_holidays_for_month(start_date, end_date)
            if temp:
                holidays = [int(i) for i in temp]

            days_in_range = (end_date - start_date).days
            days_in_range = days_in_range + 1

            working_days = days_in_range - len(holidays)
            start_date = start_date.strftime("%Y-%m-%d")
            end_date = end_date.strftime("%Y-%m-%d")

            employees = Employee().get_all_active_employees()
            results = Employee().get_att_exclude_employees_code()
            if results:
                for each_item in results:
                    c_level_emps.append(each_item.emp_code)

            att_logs = AttendanceDA().get_attendance_by_range(start_date, end_date)
            if att_logs:
                for each_item in att_logs:
                    if each_item.emp_code in att_dict:
                        temp_list = att_dict[each_item.emp_code]
                        temp_list.append(each_item)
                        att_dict[each_item.emp_code] = temp_list

                    else:
                        att_dict[each_item.emp_code] = [each_item]


            if employees:
                for emp_code in employees:
                    employee = employees[emp_code]

                    dto = att_dto()
                    dto.emp_code = emp_code
                    dto.emp_name = employee.emp_name
                    dto.gender = employee.gender
                    dto.doj = employee.doj
                    dto.designation = employee.designation
                    dto.on_time = ''
                    dto.late = ''
                    dto.late_color = "000000"
                    dto.work_hours = '-'
                    dto.work_hours_required = '-'
                    dto.work_hours_color = "000000"
                    dto.att = '-'
                    dto.leave = '-' #working_days

                    #print ('dto.emp_name', dto.emp_name)


                    dto.day_list = []

                    if emp_code in c_level_emps:
                        dto.on_time = 100
                        dto.late = 0
                        dto.work_hours = '-' #28800 * working_days #round(dto.work_hours/3600, 0)
                        dto.att = 100
                        dto.leave = 0
                    else:
                        att_data = att_dict.get(emp_code, None)

                        if att_data:
                            #print ('emp_code', emp_code)

                            tmp_dto = self.__procees_emp_attendance(att_data)

                            # if emp_code == '529':
                            #     print ('tmp_dto.work_hours', tmp_dto.work_hours)


                            dto.on_time = round((tmp_dto.days_punctual / tmp_dto.days_present) * 100, 0)
                            dto.late = round((tmp_dto.days_late / tmp_dto.days_present) * 100, 0)
                            dto.work_hours = tmp_dto.work_hours
                            dto.att = round((tmp_dto.days_present / working_days) * 100, 1)
                            dto.leave = working_days - tmp_dto.days_present

                            if dto.late > settings.THRESHOLD['LATE']:
                                dto.late_color = "FF0000"

                            #pending_time = tmp_dto.days_present -

                            dto.work_hours_required = (tmp_dto.days_present * 8) * 60 * 60
                            pending_hrs = dto.work_hours_required - dto.work_hours

                            if pending_hrs > 0:
                                if (pending_hrs/dto.work_hours_required ) * 100 > settings.THRESHOLD['WORK_HOUR']:
                                    dto.work_hours_color = "FF0000"

                            # if emp_code == '529':
                            #     print ('work_hours1111111', dto.work_hours)

                            dto.work_hours = round(dto.work_hours/3600, 0)
                            # if emp_code == '529':
                            #     print ('work_hours222222', dto.work_hours)

                            dto.work_hours_required = round(dto.work_hours_required/3600, 0)

                    dm_emp_att_list.append(dto)
                    del dto

            self.format_excel_report(dm_emp_att_list, start_date, end_date)


        except Exception as err:
            Utility().log("Error in the method generate_weekly_att_report, Error is : {0} ".format(str(err)))