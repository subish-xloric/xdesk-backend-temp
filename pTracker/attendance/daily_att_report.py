
import os
from os import path

import datetime
import time

from django.conf import settings
from types import SimpleNamespace

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side
from openpyxl.styles import Color, Fill
from openpyxl.styles import Font

from pTracker.dataaccess.db import Connection
from pTracker.common.utility import Utility
from pTracker.user_management.employee import Employee
from pTracker.dataaccess.ptracker_access.attendance import AttendanceDA
from pTracker.cronjobs.email_sender import send_daily_att_report


def att_dto():
    dto = SimpleNamespace()
    dto.emp_code = None
    dto.emp_name = None
    dto.gender = None
    dto.doj = None
    dto.designation = None
    dto.work_hours = 0
    dto.att_status = 'X'
    return dto


def new_dto():
    dto = SimpleNamespace()
    return dto


class DailyAttendanceReportBL():

    def format_excel_report(self, log_list, str_date, company_name):

        try:
            obj_date = Utility().convert_string_to_date_time(str_date, "%Y-%m-%d")
            str_date = Utility().convert_date_time_to_string(obj_date)

            month_year = Utility().convert_date_time_to_string(obj_date, "%B_%Y")
            str_long_date = Utility().convert_date_time_to_string(obj_date, "%d-%B-%Y")

            bold = Font(bold=True)
            header_fill = PatternFill(
                start_color="F0F802", end_color="F0F802", fill_type='solid')

            thin_border = Border(
                left=Side(style='thin'),
                right=Side(style='thin'),
                top=Side(style='thin'),
                bottom=Side(style='thin'))

            if log_list:
                log_list.sort(key=lambda x: x.emp_name, reverse=False)

                wb = openpyxl.Workbook()
                ws = wb.worksheets[0]
                ws.merge_cells('A1:H1')
                if company_name == "DM":
                    company_address = """Digital Mesh Softech India (P) Limited\nUnit 1: 43-A, E Block, 2nd Floor,\nCochin Special Economic Zone, Kakkanad, Kochi – 682 037, Kerala, India.\nTel: +91-484-4060200, Fax: +91-484-4060201"""
                else:
                    company_address = """EM Softech LLP\nUnit 1:Plot No.43/ A, D Block, 2nd floor,\nCochin Special Economic Zone(CSEZ), Kakkanad, Kochi-682037, Kerala, India.\nTel:+91-484-2413280"""

                ws['A1'].value = company_address
                ws['A1'].font = bold
                ws['A1'].alignment = Alignment(horizontal='center')
                #ws['A1'].fill = header_fill
                ws.row_dimensions[1].height = 95

                ws.merge_cells('A3:H3')
                ws['A3'].value = 'Attendance Report for ' + str_long_date
                ws['A3'].font = bold
                ws['A3'].alignment = Alignment(horizontal='center')
                ws['A3'].fill = header_fill

                ws['A5'].value = 'Sl No.'
                ws['A5'].font = bold

                ws['B5'].value = 'Emp Code'
                ws['B5'].font = bold

                ws['C5'].value = 'Emp Name'
                ws['C5'].font = bold

                ws['D5'].value = 'Gender'
                ws['D5'].font = bold

                ws['E5'].value = 'DOJ'
                ws['E5'].font = bold

                ws['F5'].value = 'Designation'
                ws['F5'].font = bold

                ws['G5'].value = 'Attendance'
                ws['G5'].font = bold

                ws['H5'].value = 'Work Hours'
                ws['H5'].font = bold

                #ws['I5'].value = 'Remarks'
                #ws['I5'].font = bold

                row = 6
                counter = 1
                for each_item in log_list:

                    ws['A'+str(row)].value = counter
                    ws['A'+str(row)].alignment = Alignment(horizontal='left')
                    #ws['A'+str(row)].fill = current_fill
                    #ws['A'+str(row)].border = thin_border

                    ws['B'+str(row)].value = each_item.emp_code

                    ws['C'+str(row)].value = each_item.emp_name
                    #ws['B'+str(row)].fill = current_fill
                    #ws['B'+str(row)].border = thin_border

                    ws['D'+str(row)].value = each_item.gender
                    #ws['C'+str(row)].fill = current_fill
                    #ws['C'+str(row)].border = thin_border

                    ws['E'+str(row)].value = Utility().convert_date_time_to_string(each_item.doj)
                    #ws['D'+str(row)].fill = current_fill
                    #ws['D'+str(row)].border = thin_border

                    ws['F'+str(row)].value = each_item.designation
                    #ws['E'+str(row)].fill = current_fill
                    #ws['E'+str(row)].border = thin_border

                    ws['G'+str(row)].value = each_item.att_status
                    #ws['F'+str(row)].fill = current_fill
                    #ws['F'+str(row)].border = thin_border
                    ws['G'+str(row)].alignment = Alignment(horizontal='center')
                    ws['G'+str(row)].font = bold

                    if each_item.att_status == "A":
                        ws['G'+str(row)].font = Font(color="ff0000")

                    ws['H'+str(row)].value = Utility().seconds_converter(each_item.work_hours)
                    #ws['G'+str(row)].fill = current_fill
                    #ws['G'+str(row)].border = thin_border

                    row += 1
                    counter += 1

                ws.column_dimensions["A"].width = 10
                ws.column_dimensions["B"].width = 10
                ws.column_dimensions["C"].width = 25
                ws.column_dimensions["D"].width = 12
                ws.column_dimensions["E"].width = 12

                ws.column_dimensions["F"].width = 30
                ws.column_dimensions["G"].width = 15
                ws.column_dimensions["H"].width = 15
                #ws.column_dimensions["I"].width = 15

                str_date = str_date.replace("/", "_")
                file_name = company_name + "_" + str(str_date)

                upload_path = settings.UPLOAD_PATH['DAILY_ATT_REPORT'] + month_year
                if not os.path.exists(upload_path):
                    os.makedirs(upload_path)

                file_path = upload_path + "/" + "{0}.xlsx".format(file_name)
                wb.save(file_path)


                ws.delete_cols(8, 1)
                file_name = company_name + "_" + str(str_date) + "_original"
                file_path_temp = upload_path + "/" + "{0}.xlsx".format(file_name)
                wb.save(file_path_temp)

                send_daily_att_report.apply_async([file_path, str_long_date], queue=settings.CELERY_QUEUE['mail_sender'])
        except Exception as err:
            Utility().log("Error in the method format_excel_report, Error is {0} ".format(str(err)))



    def generate_daily_att_report(self, str_date):
        att_dict = {}
        dm_emp_att_list = []
        em_emp_att_list = []
        c_level_emps = []

        try:
            employees = Employee().get_all_active_employees()
            results = Employee().get_att_exclude_employees_code()
            if results:
                for each_item in results:
                    c_level_emps.append(each_item.emp_code)

            att_logs = AttendanceDA().get_attendance_by_range(str_date, str_date)
            if att_logs:
                for each_item in att_logs:
                    att_dict[each_item.emp_code] = each_item

            if employees:
                for emp_code in employees:
                    employee = employees[emp_code]
                    dto = att_dto()
                    dto.emp_code = emp_code
                    dto.emp_name = employee.emp_name
                    dto.gender = employee.gender
                    dto.doj = employee.doj
                    dto.designation = employee.designation

                    if emp_code in c_level_emps:
                        dto.work_hours = 28800
                    else:
                        att_data = att_dict.get(emp_code, None)
                        if att_data:
                            dto.work_hours = att_data.work_hours
                        else:
                            dto.att_status = 'A'

                    if employee.company_id == settings.COMPANY['DM']['ID']:
                        dm_emp_att_list.append(dto)
                    else:
                        em_emp_att_list.append(dto)
                    del dto

            self.format_excel_report(dm_emp_att_list, str_date, 'DM')
            self.format_excel_report(em_emp_att_list, str_date, 'EM')

        except Exception as err:
            Utility().log("Error in the method generate_daily_att_report, Error is : {0} ".format(str(err)))
