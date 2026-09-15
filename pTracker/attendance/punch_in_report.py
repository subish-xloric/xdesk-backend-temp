
import os
#import datetime
from datetime import datetime, timedelta
import time

from django.conf import settings

from types import SimpleNamespace

import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side

from pTracker.dataaccess.db import Connection
from pTracker.common.utility import Utility
from pTracker.user_management.employee import Employee
from pTracker.cronjobs.email_sender import send_daily_punch_in_report

def new_dto():
    dto = SimpleNamespace()
    return dto

class PunchInBL():

    def __init__(self):
        pass

    def utc2local(self):
        return datetime.now() + timedelta(hours=5, minutes=30)

    def __process_time(self, str_date):
        str_time = None
        try:
            str_time = str(str_date).strip()
            str_time = str_time.split(" ")[1]
        except Exception as err:
            msg = "Error at __process_time, Error: {0}".format(err)
            Utility().log(msg)
        return str_time

    def get_first_punch_in(self, str_date):
        punch_in_dict = {}

        device_serial_no = settings.ATT_DEVICE['DM_IN']['SerialNumber']
        em_device_serial_no = settings.ATT_DEVICE['EM_IN']['SerialNumber']
        str_sql = """SELECT EmployeeCode, LogDate
        FROM Att
        WHERE
        (DeviceSerialNo='{0}' OR DeviceSerialNo='{2}')
        AND LogDate >= '{1} 00:00:00' AND LogDate <= '{1} 23:59:59'
        ORDER BY LogDate""".format(device_serial_no, str_date, em_device_serial_no)
        conn = Connection('essl_db')
        results, error = conn.execute(str_sql)
        if error:
            Utility().log(error)
        if results:
            for each_item in results:
                emp_code = str(each_item[0]).strip()
                if emp_code in punch_in_dict:
                    continue
                str_time = self.__process_time(each_item[1])
                punch_in_dict[emp_code] = str_time
        return punch_in_dict

    def __create_date_time(self, str_date, str_time):
        date_string = str_date + " "  + str_time
        log_time = datetime.strptime(date_string, "%Y-%m-%d %H:%M:%S")
        return log_time

    def generate_daily_punch_in_report(self, str_date):

        all_employees = []
        missed_employees = []
        c_level_emps = []
        total_employees = 0
        punctual_employees = 0
        moderate_employees = 0
        aggressive_employees = 0
        leave_employees = 0

        try:
            obj_date = Utility().convert_string_to_date_time(str_date, "%Y-%m-%d")
            month_year = Utility().convert_date_time_to_string(obj_date, "%B_%Y")
            str_long_date = Utility().convert_date_time_to_string(obj_date, "%d-%B-%Y")

            emps = Employee().get_all_active_employees()
            if emps:
                for key in emps:
                    all_employees.append(key)
                    total_employees += 1

            results = Employee().get_att_exclude_employees_code()
            if results:
                for each_item in results:
                    c_level_emps.append(each_item.emp_code)
                    total_employees = total_employees - 1

            punch_in_data = self.get_first_punch_in(str_date)

            if punch_in_data:

                punched_employees = []
                punched_employees_list = []

                punctual_time = settings.PUNCH_IN_CONFIG['punctual']['end_time']
                punctual_time = Utility().create_date_time(str_date, punctual_time)

                moderate_time = settings.PUNCH_IN_CONFIG['moderate']['end_time']
                moderate_time = Utility().create_date_time(str_date, moderate_time)

                aggressive_time = settings.PUNCH_IN_CONFIG['aggressive']['end_time']
                aggressive_time = Utility().create_date_time(str_date, aggressive_time)

                for emp_code, log_time in punch_in_data.items():
                    punched_employees.append(emp_code)
                    punch_in_time = Utility().create_date_time(str_date, log_time)

                    employee_dto = emps.get(emp_code, None)
                    if employee_dto == None:
                        continue

                    dto = new_dto()
                    dto.log_time = log_time
                    dto.emp_name = "-"
                    dto.punch_in_time = punch_in_time
                    dto.emp_code = emp_code


                    if employee_dto:
                        dto.emp_name = employee_dto.emp_name

                    if punch_in_time <= punctual_time:
                        dto.punch_type = "punctual"
                        punctual_employees += 1

                    elif punch_in_time <= moderate_time:
                        dto.punch_type = "moderate"
                        moderate_employees += 1
                    else:
                        dto.punch_type = "aggressive"
                        aggressive_employees += 1

                    punched_employees_list.append(dto)
                    del dto

                if punched_employees_list:
                    missed_employees = set(all_employees) - set(punched_employees)

                    thin_border = Border(
                        left=Side(style='thin'),
                        right=Side(style='thin'),
                        top=Side(style='thin'),
                        bottom=Side(style='thin'))


                    bold = Font(bold=True)
                    punctual_color = settings.PUNCH_IN_CONFIG['punctual']['color']
                    punctual_fill = PatternFill(start_color=punctual_color, end_color=punctual_color, fill_type='solid')

                    moderate_color = settings.PUNCH_IN_CONFIG['moderate']['color']
                    moderate_fill = PatternFill(start_color=moderate_color, end_color=moderate_color, fill_type='solid')

                    aggressive_color = settings.PUNCH_IN_CONFIG['aggressive']['color']
                    aggressive_fill = PatternFill(start_color=aggressive_color, end_color=aggressive_color, fill_type='solid')

                    missed_fill = PatternFill(start_color="39DAF4", end_color="39DAF4", fill_type='solid')
                    header_fill = PatternFill(start_color="F0F802", end_color="F0F802", fill_type='solid')


                    wb = openpyxl.Workbook()
                    ws = wb.worksheets[0]

                    ws.merge_cells('A1:J1')
                    str_dt = Utility().convert_string_to_date_time(str_date, "%Y-%m-%d")

                    # commented when time zone changed from  utc to asia/kolkota
                    # run_time = self.utc2local()

                    # comment if time changed back to utc
                    run_time = datetime.now()

                    run_time = datetime.strftime(run_time, "%I:%M %p")

                    ws['A1'].value = "Daily Punch IN Report for {0}, Run at {1} ".format(str_dt.strftime("%d %B %Y"), run_time)
                    ws['A1'].font = bold
                    ws['A1'].alignment = Alignment(horizontal='center')
                    ws['A1'].fill = header_fill


                    ws.merge_cells('A3:B3')
                    ws.merge_cells('D3:E3')
                    ws.merge_cells('G3:H3')

                    ws.merge_cells('A4:B4')
                    ws.merge_cells('D4:E4')
                    ws.merge_cells('G4:H4')

                    ws['A3'].value = 'ON TIME'
                    ws['A3'].font = bold
                    ws['A3'].alignment = Alignment(horizontal='center')
                    ws['A3'].fill = punctual_fill

                    ws['D3'].value = 'MODERATE (9:01 am - 9:30 am)'
                    ws['D3'].font = bold
                    ws['D3'].alignment = Alignment(horizontal='center')
                    ws['D3'].fill = moderate_fill

                    ws['G3'].value = 'LATE ( > 9:31 am)'
                    ws['G3'].font = bold
                    ws['G3'].alignment = Alignment(horizontal='center')
                    ws['G3'].fill = aggressive_fill

                    ws['J3'].value = 'NOT YET PUNCHED'
                    ws['J3'].font = bold
                    ws['J3'].alignment = Alignment(horizontal='center')
                    ws['J3'].fill = missed_fill

                    punctual_row = 6
                    moderate_row = 6
                    aggressive_row = 6

                    for each_emp in punched_employees_list:
                        if each_emp.punch_type == "punctual":
                            ws['A' + str(punctual_row)].value=each_emp.log_time
                            ws['B' + str(punctual_row)].value=each_emp.emp_name  #+ " (" + str(each_emp.emp_code) + ")"

                            ws['A' + str(punctual_row)].fill=punctual_fill
                            ws['B' + str(punctual_row)].fill=punctual_fill

                            ws['A'+str(punctual_row)].border=thin_border
                            ws['B'+str(punctual_row)].border=thin_border

                            punctual_row += 1
                        elif each_emp.punch_type == "moderate":
                            ws['D'+str(moderate_row)].value = each_emp.log_time
                            ws['E'+str(moderate_row)].value = each_emp.emp_name  #+ " (" + str(each_emp.emp_code) + ")"
                            ws['D'+str(moderate_row)].fill = moderate_fill
                            ws['E'+str(moderate_row)].fill = moderate_fill

                            ws['D'+str(moderate_row)].border = thin_border
                            ws['E'+str(moderate_row)].border = thin_border

                            moderate_row += 1
                        elif each_emp.punch_type == "aggressive":
                            ws['G'+str(aggressive_row)].value = each_emp.log_time
                            ws['H'+str(aggressive_row)].value = each_emp.emp_name  #+ " (" + str(each_emp.emp_code) + ")"
                            ws['G'+str(aggressive_row)].fill = aggressive_fill
                            ws['H'+str(aggressive_row)].fill = aggressive_fill
                            ws['G'+str(aggressive_row)].border = thin_border
                            ws['H'+str(aggressive_row)].border = thin_border
                            aggressive_row += 1

                    ws.column_dimensions["A"].width = 12
                    ws.column_dimensions["B"].width = 27
                    ws.column_dimensions["D"].width = 12
                    ws.column_dimensions["E"].width = 27
                    ws.column_dimensions["G"].width = 12
                    ws.column_dimensions["H"].width = 27
                    ws.column_dimensions["J"].width = 27

                if missed_employees:
                    row = 6
                    for each_code in missed_employees:
                        try:
                            temp_code = int(each_code)
                        except:
                            continue
                        if each_code in  c_level_emps:
                            continue

                        employee_dto = emps.get(each_code, None)
                        if employee_dto:
                            leave_employees += 1
                            ws['J'+str(row)].value = employee_dto.emp_name
                            ws['J'+str(row)].fill = missed_fill
                            ws['J'+str(row)].border = thin_border
                            row += 1

                ws['A4'].value = str(round((punctual_employees / total_employees ) * 100, 1)) + " %"
                ws['A4'].font = bold
                ws['A4'].alignment = Alignment(horizontal='center')
                ws['A4'].fill = punctual_fill

                ws['D4'].value = str(round((moderate_employees / total_employees ) * 100, 1)) + " %"
                ws['D4'].font = bold
                ws['D4'].alignment = Alignment(horizontal='center')
                ws['D4'].fill = moderate_fill

                ws['G4'].value = str(round((aggressive_employees / total_employees ) * 100, 1)) + " %"
                ws['G4'].font = bold
                ws['G4'].alignment = Alignment(horizontal='center')
                ws['G4'].fill = aggressive_fill

                ws['J4'].value = str(round((leave_employees / total_employees ) * 100, 1)) + " %"
                ws['J4'].font = bold
                ws['J4'].alignment = Alignment(horizontal='center')
                ws['J4'].fill = missed_fill

                upload_path = settings.UPLOAD_PATH['PUNCH_IN_REPORT'] +  month_year
                if not os.path.exists(upload_path):
                    os.makedirs(upload_path)

                str_date = Utility().convert_date_time_to_string(obj_date, "%d_%m_%Y")
                #file_name = 'punch_in_report_' + str_date
                file_path = upload_path + "/" + "{0}.xlsx".format(str_date)
                wb.save(file_path)
                send_daily_punch_in_report.apply_async([file_path, str_long_date],queue=settings.CELERY_QUEUE['mail_sender'])

                #wb.save(settings.UPLOAD_PATH['PUNCH_IN_REPORT'] + "{0}.xlsx".format(file_name))

        except Exception as err:
            Utility().log("Error in the method generate_daily_punch_in_report, Error is {0}".format(str(err)))
