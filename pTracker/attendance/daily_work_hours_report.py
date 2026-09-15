import datetime
import time
import os
from django.conf import settings
from types import SimpleNamespace
import openpyxl
from openpyxl.styles import PatternFill, Border, Side, Alignment, Protection, Font
from openpyxl.styles import Alignment
from openpyxl.styles.borders import Border, Side

from pTracker.dataaccess.db import Connection
from pTracker.common.utility import Utility
from pTracker.user_management.employee import Employee
from pTracker.dataaccess.ptracker_access.attendance import AttendanceDA
from pTracker.cronjobs.email_sender import send_daily_work_hours_report


def new_dto():
    dto = SimpleNamespace()
    return dto

def emp_dto():
    dto = SimpleNamespace()
    dto.emp_code = None
    dto.emp_name = None
    dto.arrival = None
    dto.departure = None
    dto.dm_arrival = None
    dto.dm_departure = None
    dto.em_arrival = None
    dto.em_departure = None
    dto.dm_floor_hours = 0
    dto.em_floor_hours = 0
    dto.total_floor_hours = 0
    dto.total_hours =  0
    dto.break_hours = 0
    dto.error_message = ''
    dto.arrival_status = None
    dto.departure_status = None
    return dto

class DailyWorkHoursReportBL():



    def get_daily_punch_data(self, str_date, company_code='DM', emp_code=None):
        device_serial_no_in = settings.ATT_DEVICE[company_code + "_IN"]['SerialNumber']
        device_serial_no_out = settings.ATT_DEVICE[company_code + "_OUT"]['SerialNumber']
        if emp_code:
            str_sql = """SELECT EmployeeCode, LogDate, Direction, DeviceSerialNo
                FROM Att
                WHERE
                (DeviceSerialNo='{0}' OR DeviceSerialNo='{1}')
                AND LogDate >= '{2} 00:00:00' AND LogDate <= '{2} 23:59:59'
                AND EmployeeCode = '{3}'
                ORDER BY LogDate""".format(device_serial_no_in, device_serial_no_out, str_date, emp_code)
        else:
            str_sql = """SELECT EmployeeCode, LogDate, Direction, DeviceSerialNo
                FROM Att
                WHERE
                (DeviceSerialNo='{0}' OR DeviceSerialNo='{1}')
                AND LogDate >= '{2} 00:00:00' AND LogDate <= '{2} 23:59:59'
                ORDER BY LogDate""".format(device_serial_no_in, device_serial_no_out, str_date)

        conn = Connection('essl_db')
        results, error = conn.execute(str_sql)
        if error:
            Utility().log(error)
        return results

    def __segregate_punch_data(self, punch_data):
        employee_punch_data = {}
        emps = Employee().get_all_active_employees()
        if punch_data:
            for each_item in punch_data:
                dto = new_dto()
                dto.emp_code = str(each_item[0]).strip()
                dto.emp_name = emps.get(dto.emp_code, '-')
                dto.time = Utility().convert_string_to_date_time(each_item[1])
                dto.direction = str(each_item[2]).strip().upper()
                dto.serial_no = str(each_item[3]).strip().upper()

                if dto.emp_code in employee_punch_data:
                    temp_list = employee_punch_data[dto.emp_code]
                    temp_list.append(dto)
                    employee_punch_data[dto.emp_code] = temp_list
                    del temp_list
                else:
                    employee_punch_data[dto.emp_code] = [dto]
                del dto
        return employee_punch_data

    def __process_admin_punch_data(self, emp_code, emp_name, dm_data=None, em_data=None):

        dto = emp_dto()
        dto.emp_code = emp_code
        dto.emp_name = emp_name
        log_list = []
        floor_hours = 0

        if dm_data:
            for log in dm_data:
                log_list.append(log.time)

        if em_data:
            for log in em_data:
                log_list.append(log.time)

        log_list.sort(key = lambda date: date)
        try:
            dto.dm_arrival = log_list[0]
            dto.dm_departure = log_list[-1]
            dto.arrival = log_list[0]
            dto.departure = log_list[-1]
        except:
            pass

        if dto.dm_arrival and dto.dm_departure:
            dto.total_hours = Utility().time_diff_in_seconds(dto.dm_arrival, dto.dm_departure)

        dto.total_floor_hours = dto.total_hours
        return dto

    def __process_in_and_out(self, company_code, punch_data):

        temp_missed_log = []
        j = 1
        last_direction = "-"
        floor_hours = 0
        arrival = None
        departure = None

        temp_dto = new_dto()
        temp_dto.arrival = None
        temp_dto.departure = None
        temp_dto.floor_hours = 0
        temp_dto.message = ''

        device_in = settings.ATT_DEVICE[company_code + "_IN"]['SerialNumber']
        device_out = settings.ATT_DEVICE[company_code + "_OUT"]['SerialNumber']
        number_of_log = len(punch_data)

        for log in punch_data:
            if log.serial_no == device_in:
                #departure = None
                if arrival:
                    if log.time < arrival:
                        arrival = log.time
                else:
                    arrival = log.time

                last_direction = 'IN'
                if number_of_log == j:
                    msg = """Missed the OUT for IN at {0}""".format(Utility().split_time_from_date_time(log.time))
                    temp_missed_log.append(msg)
                    break

                if punch_data[j].serial_no == device_out:
                    start_time = log.time
                    end_time = punch_data[j].time
                    floor_hours += Utility().time_diff_in_seconds(start_time, end_time)

                else:
                    msg = """Missed the OUT for IN at {0}""".format(Utility().split_time_from_date_time(log.time))
                    temp_missed_log.append(msg)
            else:
                if departure:
                    if log.time > departure:
                        departure = log.time
                else:
                    departure = log.time

                if last_direction != "IN":
                    msg = """Missed the IN for OUT at {0}""".format(Utility().split_time_from_date_time(log.time))
                    temp_missed_log.append(msg)
                last_direction = "OUT"
            j += 1

        temp_dto.arrival = arrival
        temp_dto.departure = departure
        temp_dto.floor_hours = floor_hours
        temp_dto.message = ''
        if temp_missed_log:
            temp_dto.message = ", ".join(temp_missed_log)
        return temp_dto


    def __process_punch_data(self, emp_code, emp_name, dm_data=None, em_data=None):

        dto = emp_dto()
        dto.emp_code = emp_code
        dto.emp_name = emp_name


        if dm_data:
            temp_dto = self.__process_in_and_out('DM', dm_data)
            dto.dm_arrival = temp_dto.arrival
            dto.dm_departure = temp_dto.departure
            dto.dm_floor_hours = temp_dto.floor_hours
            if temp_dto.message:
                dto.error_message = temp_dto.message
            del temp_dto

        if em_data:
            temp_dto = self.__process_in_and_out('EM', em_data)
            dto.em_arrival = temp_dto.arrival
            dto.em_departure = temp_dto.departure
            dto.em_floor_hours = temp_dto.floor_hours
            if temp_dto.message:
                dto.error_message = dto.error_message + " ." + temp_dto.message
            del temp_dto

        if dto.dm_arrival:
            dto.arrival = dto.dm_arrival

        if dto.em_arrival:
            if dto.arrival:
                if dto.em_arrival < dto.arrival:
                    dto.arrival = dto.em_arrival
            else:
                dto.arrival = dto.em_arrival

        if dto.dm_departure:
            dto.departure = dto.dm_departure

        if dto.em_departure:
            if not dto.departure:
                dto.departure = dto.em_departure
            else:
                if dto.departure < dto.em_departure:
                    dto.departure = dto.em_departure


        if dto.arrival and dto.departure:
            dto.total_hours = Utility().time_diff_in_seconds(dto.arrival, dto.departure)

        dto.total_floor_hours = dto.dm_floor_hours + dto.em_floor_hours

        if dto.total_hours:
            dto.break_hours = dto.total_hours - dto.total_floor_hours

        return dto

    def test(self):
        dates = ['2023-08-18']

        #dates = ['2019-11-23']

        for each_date in dates:
            self.generate_daily_work_hours_report(each_date)

    def generate_daily_work_hours_report(self, str_date):
        log_list = []
        emp_name = '-'
        dm_punch_data = self.get_daily_punch_data(str_date, 'DM')
        dm_emp_punch_data = self.__segregate_punch_data(dm_punch_data)

        em_punch_data = self.get_daily_punch_data(str_date, 'EM')
        em_emp_punch_data = self.__segregate_punch_data(em_punch_data)


        em_processed_list = []

        emps = Employee().get_all_active_employees()

        admin_employees = []
        admin_emps = Employee().get_admin_employees_code()
        if admin_emps:
            for each_item in admin_emps:
                admin_employees.append(each_item.emp_code)

        exclude_employees = []
        results = Employee().get_att_exclude_employees_code()
        if results:
            for each_item in results:
                exclude_employees.append(each_item.emp_code)



        if dm_emp_punch_data:
            for emp_code in dm_emp_punch_data:
                if emp_code in exclude_employees:
                    continue

                employee_dto = emps.get(emp_code, None)
                if not employee_dto:
                    msg = "Not able to find the Employee with code {0}".format(emp_code)
                    Utility().log(msg)
                    continue

                emp_name = employee_dto.emp_name
                company_id = employee_dto.company_id
                dm_data = dm_emp_punch_data[emp_code]
                em_data = em_emp_punch_data.get(emp_code, None)
                if em_data:
                    em_processed_list.append(emp_code)

                if emp_code in admin_employees:
                    dto = self.__process_admin_punch_data(emp_code, emp_name, dm_data, em_data)
                else:
                    dto = self.__process_punch_data(emp_code, emp_name, dm_data, em_data)

                dto.company_id = company_id
                dto.attendance_date = str_date
                log_list.append(dto)

                AttendanceDA().create_daily_attendance(dto)



        if em_emp_punch_data:
            for emp_code in em_emp_punch_data:
                if emp_code in exclude_employees:
                    continue

                if emp_code in em_processed_list:
                    continue

                employee_dto = emps.get(emp_code, None)
                if not employee_dto:
                    msg = "Not able to find the Employee with code {0}".format(emp_code)
                    Utility().log(msg)
                    continue

                emp_name = employee_dto.emp_name
                company_id = employee_dto.company_id
                em_data = em_emp_punch_data.get(emp_code, None)
                if em_data:
                    if emp_code in admin_employees:
                        dto = self.__process_admin_punch_data(emp_code, emp_name, None, em_data)
                    else:
                        dto = self.__process_punch_data(emp_code, emp_name, None, em_data)
                    dto.company_id = company_id
                    dto.attendance_date = str_date
                    log_list.append(dto)
                    AttendanceDA().create_daily_attendance(dto)
                    pass


        self.format_excel_file(log_list, str_date)


    def format_excel_file(self, log_list, str_date):

        punchual = 28800
        less_punchual = 25200
        avg = 21600
        very_less = 18000

        total_employees = 0
        punctual_employees = 0
        less_punctual_employees = 0
        avg_employees = 0
        very_less_employees = 0


        bold = Font(bold=True)

        header_fill = PatternFill(start_color="F0F802", end_color="F0F802", fill_type='solid')

        punchual_fill = PatternFill(start_color="e6ffcc", end_color="e6ffcc", fill_type='solid')

        less_punchual_fill = PatternFill(start_color="F6FA95", end_color="F6FA95", fill_type='solid')

        avg_fill = PatternFill(start_color="ffc299", end_color="ffc299", fill_type='solid')

        very_less_fill = PatternFill(start_color="ff704d", end_color="ff704d", fill_type='solid')

        thin_border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin'))

        if log_list:
            str_dt = Utility().convert_string_to_date_time(str_date, "%Y-%m-%d")

            log_list.sort(key=lambda x: x.total_floor_hours, reverse=True)

            wb = openpyxl.Workbook()
            ws = wb.worksheets[0]
            ws.merge_cells('A1:H1')
            ws['A1'].value = 'Work Hours Report for ' + str_dt.strftime("%d %B %Y")
            ws['A1'].font = bold
            ws['A1'].alignment = Alignment(horizontal='center')
            #ws['A1'].fill = header_fill

            ws['A3'].value = ' > 8 HRS'
            ws['A3'].font = bold
            ws['A3'].fill = punchual_fill
            ws['A3'].alignment = Alignment(horizontal='center')

            ws['A4'].value = '8-7 HRS'
            ws['A4'].font = bold
            ws['A4'].fill = less_punchual_fill
            ws['A4'].alignment = Alignment(horizontal='center')

            ws['A5'].value = '7-6 HRS'
            ws['A5'].font = bold
            ws['A5'].fill = avg_fill
            ws['A5'].alignment = Alignment(horizontal='center')

            ws['A6'].value = '< 6 HRS'
            ws['A6'].font = bold
            ws['A6'].fill = very_less_fill
            ws['A6'].alignment = Alignment(horizontal='center')


            ws['A8'].value = 'Emp Code'
            ws['A8'].font = bold

            ws['B8'].value = 'Emp Name'
            ws['B8'].font = bold

            ws['C8'].value = 'Arrival'
            ws['C8'].font = bold

            ws['D8'].value = 'Departure'
            ws['D8'].font = bold

            ws['E8'].value = 'Total Hours'
            ws['E8'].font = bold

            ws['F8'].value = 'Work Hours'
            ws['F8'].font = bold

            ws['G8'].value = 'Break Hours'
            ws['G8'].font = bold

            # ws['H3'].value = 'DM Hours'
            # ws['H3'].font = bold

            # ws['I3'].value = 'EM Hours'
            # ws['I3'].font = bold

            #ws['H8'].value = 'Remarks'
            #ws['H8'].font = bold


            row = 10
            for each_item in log_list:

                total_employees += 1

                if each_item.total_floor_hours >= punchual:
                    current_fill = punchual_fill
                    punctual_employees += 1
                elif each_item.total_floor_hours >= less_punchual:
                    current_fill = less_punchual_fill
                    less_punctual_employees += 1
                elif each_item.total_floor_hours >= avg:
                    current_fill = avg_fill
                    avg_employees += 1
                else:
                    current_fill = very_less_fill
                    very_less_employees += 1


                ws['A'+str(row)].value = each_item.emp_code
                ws['A'+str(row)].fill = current_fill
                ws['A'+str(row)].border = thin_border

                ws['B'+str(row)].value = each_item.emp_name
                ws['B'+str(row)].fill = current_fill
                ws['B'+str(row)].border = thin_border

                ws['C'+str(row)].value = Utility().split_time_from_date_time(each_item.arrival)
                ws['C'+str(row)].fill = current_fill
                ws['C'+str(row)].border = thin_border

                ws['D'+str(row)].value = Utility().split_time_from_date_time(each_item.departure)
                ws['D'+str(row)].fill = current_fill
                ws['D'+str(row)].border = thin_border

                ws['E'+str(row)].value = Utility().seconds_converter(each_item.total_hours)
                ws['E'+str(row)].fill = current_fill
                ws['E'+str(row)].border = thin_border


                ws['F'+str(row)].value = Utility().seconds_converter(each_item.total_floor_hours)
                ws['F'+str(row)].fill = current_fill
                ws['F'+str(row)].border = thin_border

                ws['G'+str(row)].value = Utility().seconds_converter(each_item.break_hours)
                ws['G'+str(row)].fill = current_fill
                ws['G'+str(row)].border = thin_border

                # ws['H'+str(row)].value = Utility().seconds_converter(each_item.dm_floor_hours)
                # ws['H'+str(row)].fill = current_fill
                # ws['H'+str(row)].border = thin_border

                # ws['I'+str(row)].value = Utility().seconds_converter(each_item.em_floor_hours)
                # ws['I'+str(row)].fill = current_fill
                # ws['I'+str(row)].border = thin_border

                #ws['H'+str(row)].value = each_item.error_message
                row += 1

            ws['B3'].value = str(round((punctual_employees / total_employees ) * 100, 1)) + " %"
            ws['B3'].font = bold
            ws['B3'].fill = punchual_fill
            ws['B3'].alignment = Alignment(horizontal='center')

            ws['B4'].value = str(round((less_punctual_employees / total_employees ) * 100, 1)) + " %"
            ws['B4'].font = bold
            ws['B4'].fill = less_punchual_fill
            ws['B4'].alignment = Alignment(horizontal='center')

            ws['B5'].value = str(round((avg_employees / total_employees ) * 100, 1)) + " %"
            ws['B5'].font = bold
            ws['B5'].fill = avg_fill
            ws['B5'].alignment = Alignment(horizontal='center')

            ws['B6'].value = str(round((very_less_employees / total_employees ) * 100, 1)) + " %"
            ws['B6'].font = bold
            ws['B6'].fill = very_less_fill
            ws['B6'].alignment = Alignment(horizontal='center')


            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 25
            ws.column_dimensions["C"].width = 12
            ws.column_dimensions["D"].width = 12
            ws.column_dimensions["E"].width = 12
            ws.column_dimensions["F"].width = 15
            ws.column_dimensions["G"].width = 15
            #ws.column_dimensions["H"].width = 35
            #ws.column_dimensions["I"].width = 15

            obj_date = Utility().convert_string_to_date_time(str_date, "%Y-%m-%d")
            str_date = Utility().convert_date_time_to_string(obj_date, "%d_%m_%Y")
            str_long_date = Utility().convert_date_time_to_string(obj_date, "%d-%B-%Y")

            month_year = Utility().convert_date_time_to_string(obj_date, "%B_%Y")
            upload_path = settings.UPLOAD_PATH['DAILY_WORK_HOUR_REPORT'] + month_year
            if not os.path.exists(upload_path):
                os.makedirs(upload_path)

            str_date = Utility().convert_date_time_to_string(obj_date, "%d_%m_%Y")

            file_path = upload_path + "/" + "{0}.xlsx".format(str_date)
            wb.save(file_path)

            send_daily_work_hours_report.apply_async([file_path, str_long_date], queue=settings.CELERY_QUEUE['mail_sender'])


    # def get_daily_work_hours(self, str_date, emp_code):
    #     log_list = []
    #     emp_name = '-'
    #     dm_punch_data = self.get_daily_punch_data(str_date, 'DM', emp_code)
    #     dm_emp_punch_data = self.__segregate_punch_data(dm_punch_data)

    #     em_punch_data = self.get_daily_punch_data(str_date, 'EM', emp_code)
    #     em_emp_punch_data = self.__segregate_punch_data(em_punch_data)


    #     em_processed_list = []

    #     # emps = Employee().get_all_active_employees()

    #     admin_employees = []
    #     admin_emps = Employee().get_admin_employees_code()
    #     if admin_emps:
    #         for each_item in admin_emps:
    #             admin_employees.append(each_item.emp_code)

    #     # exclude_employees = []
    #     # results = Employee().get_att_exclude_employees_code()
    #     # if results:
    #     #     for each_item in results:
    #     #         exclude_employees.append(each_item.emp_code)



    #     if dm_emp_punch_data:
    #         for emp_code in dm_emp_punch_data:
    #             # if emp_code in exclude_employees:
    #             #     continue

    #             # employee_dto = emps.get(emp_code, None)
    #             # if not employee_dto:
    #             #     msg = "Not able to find the Employee with code {0}".format(emp_code)
    #             #     Utility().log(msg)
    #             #     continue

    #             # emp_name = employee_dto.emp_name
    #             # company_id = employee_dto.company_id
    #             dm_data = dm_emp_punch_data[emp_code]
    #             em_data = em_emp_punch_data.get(emp_code, None)
    #             if em_data:
    #                 em_processed_list.append(emp_code)

    #             if emp_code in admin_employees:
    #                 dto = self.__process_admin_punch_data(emp_code, emp_name, dm_data, em_data)
    #             else:
    #                 dto = self.__process_punch_data(emp_code, emp_name, dm_data, em_data)

    #             dto.company_id = company_id
    #             dto.attendance_date = str_date
    #             log_list.append(dto)

    #             AttendanceDA().create_daily_attendance(dto)



    #     if em_emp_punch_data:
    #         for emp_code in em_emp_punch_data:
    #             if emp_code in exclude_employees:
    #                 continue

    #             if emp_code in em_processed_list:
    #                 continue

    #             employee_dto = emps.get(emp_code, None)
    #             if not employee_dto:
    #                 msg = "Not able to find the Employee with code {0}".format(emp_code)
    #                 Utility().log(msg)
    #                 continue

    #             emp_name = employee_dto.emp_name
    #             company_id = employee_dto.company_id
    #             em_data = em_emp_punch_data.get(emp_code, None)
    #             if em_data:
    #                 if emp_code in admin_employees:
    #                     dto = self.__process_admin_punch_data(emp_code, emp_name, None, em_data)
    #                 else:
    #                     dto = self.__process_punch_data(emp_code, emp_name, None, em_data)
    #                 dto.company_id = company_id
    #                 dto.attendance_date = str_date
    #                 log_list.append(dto)
    #                 AttendanceDA().create_daily_attendance(dto)
    #                 pass


    #     self.format_excel_file(log_list, str_date)











